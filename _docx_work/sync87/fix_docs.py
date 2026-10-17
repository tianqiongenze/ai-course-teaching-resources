# -*- coding: utf-8 -*-
"""87：按《AI技术应用基础》格式风格，改造《程序设计基础》《Python程序设计》三份文档。

- 教学方案：表0（课程名/学期/备注）、12/8 表（课时 4课时160分钟、时间双班、学时单位）
- 授课计划：表0（学期/教学方式）、表1（周别加双班时段）
- 课程标准：开课时间、学习评价、模块周次
"""
import glob
import os
import re
import sys

import docx

ROOT = r"D:/工业互联网应用2026~2027第一学期（前八周）"

CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6,
      "七": 7, "八": 8, "九": 9, "十": 10, "十一": 11, "十二": 12}

COURSES = [
    {
        "name": "程序设计基础",
        "grade": "2026级",
        "weekday": "周一",
        "n_tables": 12,
        "week_range": "第1-12周",
        "old_names": ["Python编程基础"],
        "task": "WT3.1: 工业数据采集与处理程序开发",
        "context": "3.1 程序设计基础",
        "semester": "2026-2027学年第1学期",
    },
    {
        "name": "Python程序设计",
        "grade": "2025级",
        "weekday": "周三",
        "n_tables": 8,
        "week_range": "第1-8周",
        "old_names": ["Python程序设计（项目实战）"],
        "task": "WT3.3: 工业互联网数据工程项目交付",
        "context": "3.3 Python项目实战",
        "semester": "2026-2027学年第1学期",
    },
]


def uniq(row):
    out, seen = [], set()
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def set_cell(cell, text):
    ps = cell.paragraphs
    if not ps[0].runs:
        ps[0].add_run("")
    ps[0].runs[0].text = text
    for r in ps[0].runs[1:]:
        r.text = ""
    for p in ps[1:]:
        for r in p.runs:
            r.text = ""


def para_set(par, text):
    if not par.runs:
        par.add_run("")
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def cn2int(s):
    return CN.get(s, 0)


def fix_lesson_plan(cfg, apply):
    p = glob.glob(os.path.join(ROOT, cfg["grade"], "《%s》" % cfg["name"],
                               "《%s》-教学设计" % cfg["name"], "*.docx"))[0]
    d = docx.Document(p)
    log = []

    # 表0
    t0 = d.tables[0]
    rows = {uniq(r)[0].text.strip(): uniq(r) for r in t0.rows}
    # 课程名称
    for k, u in rows.items():
        if k == "课程名称" and len(u) > 1:
            if u[1].text.strip() != cfg["name"]:
                log.append(("课程名称", u[1].text.strip(), cfg["name"]))
                set_cell(u[1], cfg["name"])
        if k == "授课学期" and len(u) > 1:
            if u[1].text.strip() != cfg["semester"]:
                log.append(("授课学期", u[1].text.strip(), cfg["semester"]))
                set_cell(u[1], cfg["semester"])
        if k.startswith("备注"):
            cell = u[1]
            for par in cell.paragraphs:
                t = par.text
                if t.startswith("工作情境"):
                    new = ("工作情境：%s（%s，每周4课时，1班%s1~4节/2班5~8节；前置：%s）"
                           % (cfg["context"], cfg["week_range"], cfg["weekday"],
                              "Python编程基础前6周" if cfg["name"] == "程序设计基础" else "C语言编程基础"))
                    if new != t:
                        log.append(("备注-情境", t[:36], new[:44]))
                        para_set(par, new)
                elif t.startswith("课时与考核"):
                    new = ("课时与考核：每课时40分钟（每周4课时160分钟＝上/下2次课各80分钟）；"
                           "过程性评价80%（工单40%+课堂表现10%+小组项目30%）+期末20%"
                           "（详见课程标准与授课计划表）")
                    if new != t:
                        log.append(("备注-考核", t[:36], new[:44]))
                        para_set(par, new)

    # 课次表
    for ti in range(1, cfg["n_tables"] + 1):
        t = d.tables[ti]
        u1 = uniq(t.rows[1])
        # 课时：'4' → '4课时（160分钟）'
        for c in u1:
            if c.text.strip() == "4":
                set_cell(c, "4课时（160分钟）")
                log.append(("表%d课时" % ti, "4", "4课时（160分钟）"))
        # 时间
        for c in u1:
            txt = c.text.strip()
            m = re.search(r"第([一二三四五六七八九十]+)周星期([一二三四五六日])", txt)
            if m:
                n = cn2int(m.group(1))
                wd = cfg["weekday"]
                new = "第%d周%s1~4节（1班）、5~8节（2班）" % (n, wd)
                if new != txt:
                    set_cell(c, new)
                    log.append(("表%d时间" % ti, txt[:30], new))
        # 任务实施学时单位
        for ri in (13, 14, 16, 17, 19, 20):
            u = uniq(t.rows[ri])
            if len(u) < 5:
                continue
            h = u[4].text.strip()
            if h.endswith("min"):
                set_cell(u[4], h.replace("min", "分钟"))
                log.append(("表%dR%d学时" % (ti, ri), h, h.replace("min", "分钟")))

    for tag, old, new in log:
        print("  %-12s %-40s → %s" % (tag, old, new))
    if apply:
        d.save(p)
    print("  教学方案：%d 处%s" % (len(log), "（已保存）" if apply else "（空跑）"))


def fix_plan_table(cfg, apply):
    p = glob.glob(os.path.join(ROOT, cfg["grade"], "《%s》" % cfg["name"],
                               "《%s》-授课计划" % cfg["name"], "*.docx"))[0]
    d = docx.Document(p)
    log = []
    t0 = d.tables[0]
    for row in t0.rows:
        u = uniq(row)
        if len(u) < 2:
            continue
        k, c = u[0].text.strip(), u[1]
        if k == "授课学期" and c.text.strip() != cfg["semester"]:
            log.append(("授课学期", c.text.strip(), cfg["semester"]))
            set_cell(c, cfg["semester"])
        if k == "教学方式":
            txt = c.text
            slot = "1班%s1~4节/2班5~8节" % cfg["weekday"]
            if slot not in txt:
                if txt.startswith("每课时40分钟、每周4课时"):
                    new = txt.replace("每周4课时（160分钟）", "每周4课时（160分钟，%s）" % slot, 1)
                    if new == txt:
                        new = "每课时40分钟、每周4课时（160分钟，%s）；" % slot + txt
                else:
                    new = "每课时40分钟、每周4课时（160分钟＝上/下2次课各80分钟，%s）；" % slot + txt
                log.append(("教学方式", txt[:34], new[:44]))
                set_cell(c, new)

    # 周别
    t1 = d.tables[1]
    for ri in range(2, len(t1.rows) - 3):
        cell = t1.rows[ri].cells[0]
        txt = cell.text.strip()
        m = re.search(r"第([一二三四五六七八九十]+)周", txt)
        if not m:
            continue
        n = cn2int(m.group(1))
        new = "第%d周\n%s1~4节（1班）\n5~8节（2班）" % (n, cfg["weekday"])
        if new != txt:
            set_cell(cell, new)
            log.append(("周别R%d" % ri, txt.replace("\n", " "), new.replace("\n", " ")))

    for tag, old, new in log:
        print("  %-12s %-40s → %s" % (tag, old, new))
    if apply:
        d.save(p)
    print("  授课计划：%d 处%s" % (len(log), "（已保存）" if apply else "（空跑）"))


def fix_kebiao(cfg, apply):
    p = glob.glob(os.path.join(ROOT, cfg["grade"], "《%s》" % cfg["name"],
                               "《%s》-课程标准" % cfg["name"], "*.docx"))[0]
    d = docx.Document(p)
    log = []
    # 开课时间
    t0 = d.tables[0]
    for row in t0.rows:
        u = uniq(row)
        for c in u:
            if "开课时间" == c.text.strip():
                continue
        for i, c in enumerate(u):
            if c.text.strip().startswith(cfg["week_range"]):
                old = c.text.strip()
                new = "%s·每周4课时（1班%s1~4节/2班5~8节）" % (cfg["week_range"], cfg["weekday"])
                if new != old:
                    set_cell(c, new)
                    log.append(("开课时间", old[:34], new[:46]))
    # 学习评价
    for par in d.paragraphs:
        t = par.text.strip()
        if t.startswith("综合成绩=平时成绩"):
            new = ("综合成绩=过程性评价（80%）+期末考试（20%）；过程性评价=工单任务（40%）"
                   "+课堂表现（10%）+小组项目（30%）。")
            para_set(par, new)
            log.append(("学习评价", t[:34], new[:40]))
        elif t.startswith("过程性评价：工单完成度"):
            new = ("过程性评价（80%）：工单任务40%（8工单各5%，及格/优秀/超优三档）、"
                   "课堂表现10%（出勤·参与·互动·答问）、小组项目30%（综合项目交付与答辩）；"
                   "期末考试（20%）：上机操作与答辩。")
            para_set(par, new)
            log.append(("学习评价", t[:34], new[:40]))
    # 模块周次：补「课时」单位
    for t in d.tables:
        if len(t.rows) < 2 or "周次" not in " ".join(c.text for c in t.rows[0].cells):
            continue
        for ri in range(1, len(t.rows)):
            u = uniq(t.rows[ri])
            if len(u) < 2:
                continue
            c = u[1]
            txt = c.text.strip()
            m = re.match(r"^(第[\d\-]+周)\s*/\s*(\d+)$", txt)
            if m:
                new = "%s / %s课时" % (m.group(1), m.group(2))
                set_cell(c, new)
                log.append(("模块周次R%d" % ri, txt, new))

    for tag, old, new in log:
        print("  %-12s %-40s → %s" % (tag, old, new))
    if apply:
        d.save(p)
    print("  课程标准：%d 处%s" % (len(log), "（已保存）" if apply else "（空跑）"))


def main():
    apply = "--apply" in sys.argv
    for cfg in COURSES:
        print("\n" + "=" * 66)
        print("### %s（%s）" % (cfg["name"], cfg["grade"]))
        fix_lesson_plan(cfg, apply)
        fix_plan_table(cfg, apply)
        fix_kebiao(cfg, apply)


if __name__ == "__main__":
    main()
