# -*- coding: utf-8 -*-
"""87b：工单课时统一 + PPT 课程名统一（两门课程）。"""
import glob
import os
import re
import sys

import docx
from pptx import Presentation

ROOT = r"D:/工业互联网应用2026~2027第一学期（前八周）"

COURSES = [
    {"name": "程序设计基础", "grade": "2026级", "old": "Python编程基础"},
    {"name": "Python程序设计", "grade": "2025级", "old": "Python程序设计（项目实战）"},
]


def uniq(row):
    out, seen = [], set()
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def para_replace(par, old, new):
    if old not in par.text:
        return 0
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return 1
    if not par.runs:
        return 0
    par.runs[0].text = par.text.replace(old, new)
    for r in par.runs[1:]:
        r.text = ""
    return 1


def fix_gongdan(cfg, apply):
    D = os.path.join(ROOT, cfg["grade"], "《%s》" % cfg["name"], "《%s》-任务单" % cfg["name"])
    fs = sorted(f for f in glob.glob(os.path.join(D, "*.docx"))
                if not os.path.basename(f).startswith("~$"))
    total = 0
    for f in fs:
        d = docx.Document(f)
        n = 0
        # 头部段落
        for par in d.paragraphs:
            if "2课时（80分钟）" in par.text:
                n += para_replace(par, "2课时（80分钟）", "4课时（160分钟）")
        # 表格
        for t in d.tables:
            for row in t.rows:
                for c in uniq(row):
                    for par in c.paragraphs:
                        if "周次/课时" in par.text and "2课时" in par.text:
                            n += para_replace(par, "/2课时", "/4课时")
                        elif "2课时（80分钟）" in par.text:
                            n += para_replace(par, "2课时（80分钟）", "4课时（160分钟）")
        if n and apply:
            d.save(f)
        total += n
    print("  工单 %d 个 / %d 处%s" % (len(fs), total, "（已保存）" if apply else "（空跑）"))


def fix_ppt(cfg, apply):
    D = os.path.join(ROOT, cfg["grade"], "《%s》" % cfg["name"], "《%s》-教学PPT" % cfg["name"])
    fs = sorted(f for f in glob.glob(os.path.join(D, "*.pptx"))
                if not os.path.basename(f).startswith("~$"))
    total = 0
    for f in fs:
        prs = Presentation(f)
        n = 0
        for s in prs.slides:
            for sh in s.shapes:
                if not sh.has_text_frame:
                    continue
                for par in sh.text_frame.paragraphs:
                    if cfg["old"] in par.text:
                        n += para_replace(par, cfg["old"], cfg["name"])
        if n and apply:
            prs.save(f)
        total += n
    print("  PPT %d 册 / %d 处%s" % (len(fs), total, "（已保存）" if apply else "（空跑）"))


def main():
    apply = "--apply" in sys.argv
    for cfg in COURSES:
        print("\n### %s（%s）课程名 %s → %s" % (cfg["name"], cfg["grade"], cfg["old"], cfg["name"]))
        fix_gongdan(cfg, apply)
        fix_ppt(cfg, apply)


if __name__ == "__main__":
    main()
