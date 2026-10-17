# -*- coding: utf-8 -*-
"""84号轮：目标目录《AI技术应用基础》全量验收。"""
import glob
import os
import re

import docx
from pptx import Presentation

TGT = r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》"


def uniq(row):
    out, seen = [], set()
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def main():
    ok = True

    # ① 教学方案
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-教学设计", "*.docx"))[0]
    d = docx.Document(p)
    print("=== ① 教学方案 %s ===" % os.path.basename(p))
    print("    段落 %d / 表格 %d" % (len(d.paragraphs), len(d.tables)))
    bad = []
    for ti in range(1, 13):
        t = d.tables[ti]
        u1 = uniq(t.rows[1])
        slot = u1[-1].text.strip() if u1 else ""
        tot = 0
        for ri in (13, 14, 16, 17, 19, 20):
            u = uniq(t.rows[ri])
            m = re.match(r"(\d+)分钟", u[4].text.strip())
            if m:
                tot += int(m.group(1))
        if not re.search(r"第\d{1,2}周周五[15]~[48]节课", slot):
            bad.append("表%d 时间=%s" % (ti, slot))
        if tot != 160:
            bad.append("表%d 合计=%d分钟" % (ti, tot))
        print("    表%-2d %-18s 合计%3d分钟 %s" % (ti, slot, tot, "✔" if tot == 160 else "✘"))
    if bad:
        ok = False
        print("    ✘ 异常:", bad)

    # ② 授课计划
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-授课计划", "*.docx"))[0]
    d = docx.Document(p)
    print("\n=== ② 授课计划 %s ===" % os.path.basename(p))
    t0, t1 = d.tables[0], d.tables[1]
    print("    课程名称:", t0.rows[2].cells[1].text)
    print("    授课学期:", t0.rows[4].cells[1].text)
    print("    班级    :", t0.rows[5].cells[1].text)
    wk = [t1.rows[i].cells[0].text.replace("\n", " ") for i in range(2, 14)]
    print("    周别    :", " / ".join(wk))
    if not all("周五" in w for w in wk):
        ok = False
        print("    ✘ 周别未含周五时段")

    # ③ 课程标准
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-课程标准", "*.docx"))[0]
    d = docx.Document(p)
    print("\n=== ③ 课程标准 %s ===" % os.path.basename(p))
    print("    开课时间:", d.tables[0].rows[5].cells[3].text)
    print("    模块周次:", " | ".join(d.tables[3].rows[i].cells[1].text for i in range(1, 7)))

    # ④ 任务单
    fs = sorted(glob.glob(os.path.join(TGT, "《AI技术应用基础》-任务单", "*.docx")))
    print("\n=== ④ 任务单 %d 个 ===" % len(fs))
    for f in fs:
        d = docx.Document(f)
        hdr = [x.text for x in d.paragraphs[:4] if x.text.strip()]
        line = hdr[2] if len(hdr) > 2 else ""
        if "4课时（160分钟）" not in line:
            ok = False
            print("    ✘ %s: %s" % (os.path.basename(f)[:40], line[:50]))
    print("    全部 4课时（160分钟）标注 ✔" if ok else "")

    # ⑤ PPT
    fs = sorted(glob.glob(os.path.join(TGT, "《AI技术应用基础》-教学PPT", "*.pptx")))
    print("\n=== ⑤ 教学PPT %d 册 ===" % len(fs))
    pages = []
    for f in fs:
        try:
            n = len(Presentation(f).slides)
            pages.append(n)
        except Exception as e:
            ok = False
            print("    ✘ 打开失败 %s: %s" % (os.path.basename(f), e))
    print("    页数范围: %d ~ %d，合计 %d 页" % (min(pages), max(pages), sum(pages)))

    print("\n=== 结论:", "全部通过 ✔" if ok else "存在异常 ✘", "===")


if __name__ == "__main__":
    main()
