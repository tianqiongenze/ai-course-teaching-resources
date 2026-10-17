# -*- coding: utf-8 -*-
"""85b：时间统一为「第X周周五1~4节课（1班）、周五5~8节课（2班）」。

覆盖：教学方案 12 表「时间」、授课计划 12 周别、课标开课时间、两处备注。
"""
import glob
import os
import re
import sys

import docx

TGT = r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》"


def uniq(row):
    out, seen = [], set()
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        out.append(c)
    return out


def set_cell(cell, text):
    ps = cell.paragraphs
    first = ps[0]
    if not first.runs:
        first.add_run("")
    first.runs[0].text = text
    for r in first.runs[1:]:
        r.text = ""
    for p in ps[1:]:
        for r in p.runs:
            r.text = ""


def para_set(par, text):
    if not par.runs:
        par.add_run("")
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def main():
    apply = "--apply" in sys.argv
    log = []

    # ① 教学方案 12 表时间
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-教学设计", "*.docx"))[0]
    d = docx.Document(p)
    for ti in range(1, 13):
        u1 = uniq(d.tables[ti].rows[1])
        for i, c in enumerate(u1):
            if c.text.strip() == "时间" and i + 1 < len(u1):
                cur = u1[i + 1].text.strip()
                m = re.search(r"第(\d{1,2})周", cur)
                if m:
                    new = "第%s周周五1~4节课（1班）、周五5~8节课（2班）" % m.group(1)
                    if new != cur:
                        log.append(("教学方案表%d" % ti, cur, new))
                        set_cell(u1[i + 1], new)
    # 备注同步
    for par in d.tables[0].rows[8].cells[1].paragraphs:
        if "每周五4课时连排" in par.text:
            new = par.text.replace("每周五4课时连排", "每周五4课时连排（1班1~4节/2班5~8节）")
            log.append(("教学方案备注", par.text[:40], new[:50]))
            para_set(par, new)
        if "（每周五4课时160分钟连排）" in par.text:
            new = par.text.replace("（每周五4课时160分钟连排）", "（每周五4课时160分钟连排，1班1~4节/2班5~8节）")
            log.append(("教学方案课时备注", par.text[:40], new[:50]))
            para_set(par, new)
    if apply:
        d.save(p)

    # ② 授课计划 12 周别 + 教学方式
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-授课计划", "*.docx"))[0]
    d = docx.Document(p)
    t1 = d.tables[1]
    for i in range(12):
        cell = t1.rows[2 + i].cells[0]
        cur = cell.text.strip()
        m = re.search(r"第(\d{1,2})周", cur)
        if m:
            new = "第%s周\n周五1~4节（1班）\n周五5~8节（2班）" % m.group(1)
            if new != cur:
                log.append(("授课计划R%d" % (2 + i), cur.replace("\n", " "), new.replace("\n", " ")))
                set_cell(cell, new)
    c = d.tables[0].rows[6].cells[1]
    if "（160分钟连排" in c.text and "1班" not in c.text:
        new = c.text.replace("（160分钟连排：周五1~4节或周五5~8节）",
                             "（160分钟连排：1班周五1~4节 / 2班周五5~8节）")
        log.append(("授课计划教学方式", c.text[:36], new[:40]))
        set_cell(c, new)
    if apply:
        d.save(p)

    # ③ 课标开课时间
    p = glob.glob(os.path.join(TGT, "《AI技术应用基础》-课程标准", "*.docx"))[0]
    d = docx.Document(p)
    for row in d.tables[0].rows:
        for c in row.cells:
            if "每周五4课时" in c.text and "1班" not in c.text:
                new = c.text.replace("每周五4课时", "每周五4课时（1班1~4节/2班5~8节）")
                log.append(("课标开课时间", c.text[:40], new[:50]))
                set_cell(c, new)
    if apply:
        d.save(p)

    for tag, old, new in log:
        print("  %-14s %-44s → %s" % (tag, old, new))
    print("\n合计 %d 处；%s" % (len(log), "已写入" if apply else "空跑"))


if __name__ == "__main__":
    main()
