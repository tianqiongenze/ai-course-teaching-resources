# -*- coding: utf-8 -*-
"""84号轮：更新目标目录《AI技术应用基础》分层工单的周次与课时标注。

规则：
- 周次以文件名 W{NN} 为准（W07→第7周 … W20→第20周）
- 课时统一为「4课时（160分钟）」（原 2课时80分钟 → 每周五4课时连排）
- 表格首行「周次/课时: 第X周/2课时」→「第X周/4课时」
"""
import glob
import os
import re
import sys

import docx

TGT = r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》/《AI技术应用基础》-任务单"


def para_replace(par, old, new):
    """段落内替换，尽量保留格式：单 run 命中则原地改，否则合并到首 run。"""
    txt = par.text
    if old not in txt:
        return 0
    # 先尝试单 run 替换
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return 1
    # 跨 run：整段重写（保留首个 run 的格式）
    if not par.runs:
        return 0
    par.runs[0].text = txt.replace(old, new)
    for r in par.runs[1:]:
        r.text = ""
    return 1


def iter_paras(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def main():
    apply = "--apply" in sys.argv
    files = sorted(glob.glob(os.path.join(TGT, "*.docx")))
    print("工单数:", len(files))
    total = 0
    for f in files:
        base = os.path.basename(f)
        m = re.match(r"W(\d{2})_", base)
        if not m:
            print("  [SKIP] 无周次前缀:", base)
            continue
        wk = int(m.group(1))
        doc = docx.Document(f)
        n = 0
        for par in iter_paras(doc):
            t = par.text
            if not t:
                continue
            # ① 段落「第X周 · 2课时（80分钟）」→「第N周 · 4课时（160分钟）」
            m2 = re.search(r"第(\d{1,2})周\s*[·・]\s*2课时（80分钟）", t)
            if m2:
                n += para_replace(par, m2.group(0), "第%d周 · 4课时（160分钟）" % wk)
                continue
            # ② 表格「周次/课时: 第X周/2课时」
            m3 = re.search(r"周次/课时[:：]\s*第(\d{1,2})周/2课时", t)
            if m3:
                n += para_replace(par, m3.group(0), "周次/课时: 第%d周/4课时" % wk)
                continue
            # ③ 兜底：残留的「第X周/2课时」
            m4 = re.search(r"第(\d{1,2})周/2课时", t)
            if m4:
                n += para_replace(par, m4.group(0), "第%d周/4课时" % wk)
                continue
            # ④ 兜底：单独「2课时（80分钟）」
            if "2课时（80分钟）" in t:
                n += para_replace(par, "2课时（80分钟）", "4课时（160分钟）")
        if apply and n:
            doc.save(f)
        total += n
        print("  %-52s W%-2d 替换 %d 处%s" % (base[:52], wk, n, "" if apply else "（空跑）"))
    print("\n合计 %d 处；%s" % (total, "已写入" if apply else "空跑，加 --apply 生效"))


if __name__ == "__main__":
    main()
