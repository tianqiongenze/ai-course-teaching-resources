# -*- coding: utf-8 -*-
"""84号轮：更新《AI技术应用基础》授课计划表至 05 最新口径。

- 课程名称/学期/班级/教学方式/考核方式按 05 最新同步
- 周别列：教学周 1-12 → 实际学期周 7/10/11…20 + 周五时段（前6次1~4节、后6次5~8节）
"""
import os
import sys

import docx

DOC = (r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》"
       r"/《AI技术应用基础》-授课计划/工业互联网应用专业《AI技术应用基础》授课计划表v1.0.docx")

# 教学周 → (学期周, 周五时段)
WEEKMAP = [
    (7, "周五1~4节"), (10, "周五1~4节"), (11, "周五1~4节"), (12, "周五1~4节"),
    (13, "周五1~4节"), (14, "周五1~4节"), (15, "周五5~8节"), (16, "周五5~8节"),
    (17, "周五5~8节"), (18, "周五5~8节"), (19, "周五5~8节"), (20, "周五5~8节"),
]

TEACH_WAY = ("每课时40分钟、每周五4课时（160分钟连排：周五1~4节或周五5~8节）；"
             "理实一体'讲演做'+完整六步法（资讯·计划课前线上完成）；"
             "虚拟产线数据集+本地工具链（CPU优先，GPU可选）；ABCD分层+费曼师徒制")

ASSESS = ("过程性评价（80%）：平时工单40%（10次工单×4分，正确性+规范+留痕）、"
          "课堂表现10%（出勤·参与·互动游戏·答问）、"
          "小组项目30%（期末项目册20选1：代码+README+答辩）。\n"
          "期末考试（20%）：第20周统一组织。\n"
          "及格线：各周工单达D级及以上；综合成绩60分。")


def set_cell_text(cell, new_text):
    """写入单元格：保留首段格式，多行用换行写入同一段。"""
    paras = cell.paragraphs
    first = paras[0]
    if not first.runs:
        first.add_run("")
    first.runs[0].text = new_text
    for r in first.runs[1:]:
        r.text = ""
    for p in paras[1:]:
        for r in p.runs:
            r.text = ""


def main():
    apply = "--apply" in sys.argv
    d = docx.Document(DOC)
    t0, t1 = d.tables[0], d.tables[1]
    changes = []

    # ── 表0 基本信息 ──
    for ri, newv in [
        (2, "AI技术应用基础（48课时版）"),
        (4, "2026-2027学年第1学期"),
        (5, "工业互联网应用专业2026级"),
        (6, TEACH_WAY),
        (7, ASSESS),
    ]:
        cell = t0.rows[ri].cells[1]
        changes.append((ri, cell.text.replace("\n", " / ")[:40], newv.replace("\n", " / ")[:40]))
        set_cell_text(cell, newv)

    # ── 表1 周别列（R2..R13）──
    for i, (wk, slot) in enumerate(WEEKMAP):
        ri = 2 + i
        cell = t1.rows[ri].cells[0]
        old = cell.text.strip()
        newv = "第%d周\n%s" % (wk, slot)
        changes.append(("周别R%d" % ri, old[:16], newv.replace("\n", " / ")))
        set_cell_text(cell, newv)

    for tag, old, new in changes:
        print("  %-10s %-42s → %s" % (tag, old, new))
    if apply:
        d.save(DOC)
        print("\n已写入:", os.path.basename(DOC))
    else:
        print("\n（空跑，加 --apply 生效）")


if __name__ == "__main__":
    main()
