# -*- coding: utf-8 -*-
"""85d：修复 V2 授课 PPT 的周次污染（79 轮连锁替换 bug）。

只改两处：
  ① 封面「第X周 · 第N次课（2课时 · 80分钟）」→ 本册周次
  ② 结尾「下次课预告（第X周）」→ 下一册周次
不动「本课(第7周起·跳过W08/W09·至第20周)」等全局描述。
"""
import glob
import os
import re
import sys

from pptx import Presentation

DIRS = [
    r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》/《AI技术应用基础》-教学PPT",
    r"D:/Python编程实战/教学实施资料/05-AI应用基础48课时版/V2授课PPT",
]


def key(f):
    m = re.match(r"W(\d{2})-第(\d+)次课", os.path.basename(f))
    return (int(m.group(1)), int(m.group(2))) if m else (999, 999)


def fix_runs(tf, pattern, repl):
    """在文本框内按 run 级替换，保留格式。"""
    n = 0
    for par in tf.paragraphs:
        for r in par.runs:
            if re.search(pattern, r.text):
                new = re.sub(pattern, repl, r.text)
                if new != r.text:
                    r.text = new
                    n += 1
    return n


def main():
    apply = "--apply" in sys.argv
    total = 0
    for D in DIRS:
        fs = sorted((f for f in glob.glob(os.path.join(D, "*.pptx"))
                     if not os.path.basename(f).startswith("~$")), key=key)
        if not fs:
            print("[SKIP] 无文件:", D)
            continue
        print("\n### %s（%d 册）" % (os.path.basename(D), len(fs)))
        for i, f in enumerate(fs):
            base = os.path.basename(f)
            wk, nth = key(f)
            nxt = key(fs[i + 1])[0] if i + 1 < len(fs) else wk
            prs = Presentation(f)
            n = 0
            for s in prs.slides:
                for sh in s.shapes:
                    if not sh.has_text_frame:
                        continue
                    t = sh.text_frame.text
                    # ① 封面
                    if "次课（2课时" in t and re.search(r"第\s*\d{1,2}\s*周", t):
                        n += fix_runs(sh.text_frame, r"第\s*\d{1,2}\s*周", "第%d周" % wk)
                    # ② 下次课预告
                    elif "下次课预告" in t:
                        n += fix_runs(sh.text_frame, r"（第\s*\d{1,2}\s*周）", "（第%d周）" % nxt)
            if n and apply:
                prs.save(f)
            total += n
            print("  %-44s 本册W%-2d 预告W%-2d 修正 %d 处%s"
                  % (base[:44], wk, nxt, n, "" if apply else "（空跑）"))
    print("\n合计 %d 处；%s" % (total, "已写入" if apply else "空跑"))


if __name__ == "__main__":
    main()
