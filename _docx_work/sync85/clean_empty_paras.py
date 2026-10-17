# -*- coding: utf-8 -*-
"""85g：清理 docx 表格单元格的尾部空段落（替换文本后残留，导致行高膨胀）。

只删「尾部连续空段落」，保留至少 1 段；不动中间空段（可能是有意留白）。
"""
import glob
import os
import sys

import docx

TGT = r"D:/工业互联网应用2026~2027第一学期（前八周）/2026级/《AI技术应用基础》"
DIRS = [
    "《AI技术应用基础》-教学设计",
    "《AI技术应用基础》-授课计划",
    "《AI技术应用基础》-课程标准",
    "《AI技术应用基础》-任务单",
]


def clean_cell(cell):
    ps = cell.paragraphs
    if len(ps) <= 1:
        return 0
    # 从后往前，遇到第一个非空段落即停
    removed = 0
    for p in reversed(ps[1:]):
        if p.text.strip():
            break
        el = p._element
        el.getparent().remove(el)
        removed += 1
    return removed


def main():
    apply = "--apply" in sys.argv
    total = 0
    for sub in DIRS:
        fs = sorted(f for f in glob.glob(os.path.join(TGT, sub, "*.docx"))
                    if not os.path.basename(f).startswith("~$"))
        for f in fs:
            d = docx.Document(f)
            n = 0
            for t in d.tables:
                for row in t.rows:
                    seen = set()
                    for c in row.cells:
                        k = id(c._tc)
                        if k in seen:
                            continue
                        seen.add(k)
                        n += clean_cell(c)
            if n:
                total += n
                print("  %-46s 清理 %d 段" % (os.path.basename(f)[:46], n))
                if apply:
                    d.save(f)
    print("\n合计 %d 段；%s" % (total, "已写入" if apply else "空跑"))


if __name__ == "__main__":
    main()
