# -*- coding: utf-8 -*-
"""05包(AI48) 效果图注入（73号轮）——为 V1/V2 共48册新建「工单完成效果图」页

背景：05包 V1/V2 原本**没有**效果图页（既无占位页也无内容），
      与 01/02/03 包（均有真实运行效果图）存在一致性缺口。
做法：克隆本册标准模板页 → 仅保留页眉/页脚 → 重建三卡内容 → 插到"工单页"之后 → 页脚重编。
数据：`教师参考/solutions.py` 12 个完成态实现的**真实运行输出**（非编造）。

用法:
  python inject_shots_ai48.py --run      # 跑12工单解，缓存输出
  python inject_shots_ai48.py --inject   # 幂等注入（已有"工单完成效果图"页自动跳过）
  python inject_shots_ai48.py --dry      # 只打印计划
"""
import argparse
import contextlib
import copy
import glob
import io
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = r"D:/Python编程实战/教学实施资料"
PKG = os.path.join(ROOT, "05-AI应用基础48课时版")
CACHE = os.path.join(ROOT, "_docx_work", "slim", "_shots_ai48.json")

# 文件名关键词 → solutions 函数
KEYMAP = [
    ("数据清洗与质量评估", "m11a"),
    ("特征设计与优化", "m11b"),
    ("传统分类", "m21a"),
    ("XGBoost参数调优", "m22"),
    ("MLP神经网络", "m23a"),
    ("CNN或LSTM", "m23b"),
    ("ResNet迁移学习", "m31"),
    ("YOLO实时缺陷检测", "m32"),
    ("LLM故障诊断", "m41"),
    ("RAG知识库问答", "m42"),
    ("多智能体检修调度", "m51"),
    ("综合项目答辩", "z"),
]
FUNCS = [f for _, f in KEYMAP]

# 模板页识别：只保留页眉/页脚（按 left,top 英寸，保留2位小数）
TPL = {
    "v1": {(0.0, 0.0), (0.52, 0.13), (0.52, 7.08), (0.52, 7.06), (11.31, 7.06)},
    "v2": {(0.0, 0.0), (0.0, 1.26), (0.77, 0.33), (0.55, 7.11), (0.55, 7.15), (11.59, 7.15)},
}
# 配色（取自本册实际主题）
C = {
    "v1": dict(title="#231D3D", head="#3D2E8C", accent="#00A6A6", body="#231D3D",
               muted="#5F5A7D", card="#FFFFFF", card2="#EAE7F8", border="#CFC8E8"),
    "v2": dict(title="#FFFFFF", head="#0545E8", accent="#F5670F", body="#010328",
               muted="#44546A", card="#F3F6FD", card2="#FFFFFF", border="#C9D4EE"),
}


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def run_solutions():
    sys.path.insert(0, os.path.join(PKG, "教师参考"))
    cwd = os.getcwd()
    os.chdir(PKG)
    out = {}
    try:
        for m in list(sys.modules):
            if m == "solutions":
                del sys.modules[m]
        import solutions
        for f in FUNCS:
            buf = io.StringIO()
            try:
                with contextlib.redirect_stdout(buf):
                    getattr(solutions, f)()
                out[f] = buf.getvalue()
            except Exception as e:
                out[f] = f"__ERR__ {type(e).__name__}: {str(e)[:100]}"
    finally:
        os.chdir(cwd)
    json.dump(out, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for f in FUNCS:
        t = out[f]
        print(f"  {f:<6} {'ERR' if t.startswith('__ERR__') else str(len(t.splitlines())) + '行'}  "
              f"{t.splitlines()[0][:56] if t.strip() else ''}")
    return out


def clean_lines(text, maxlines=9, maxw=64):
    ls = [l.rstrip() for l in text.splitlines() if l.strip()]
    ls = [l if len(l) <= maxw else l[:maxw - 1] + "…" for l in ls]
    merged = []
    for l in ls:
        k = re.sub(r"\d+", "#", l)
        if merged and re.sub(r"\d+", "#", merged[-1][0]) == k:
            merged[-1][1] += 1
            continue
        merged.append([l, 1])
    res = [l + (f"  ×{c}" if c > 1 else "") for l, c in merged]
    if len(res) > maxlines:
        res = res[:maxlines - 1] + [res[-1]]
    return res


def clone_slide(prs, src_idx):
    src = prs.slides[src_idx]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(shp._element))
    for rel in src.part.rels.values():
        try:
            if rel.is_external:
                dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            else:
                dst.part.rels.get_or_add(rel.reltype, rel._target)
        except Exception:
            pass
    return dst


def move_slide(prs, from_idx, to_idx):
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[from_idx]
    lst.remove(el)
    lst.insert(to_idx, el)


def add_card(slide, x, y, w, h, fill, border):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(border)
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def add_text(slide, x, y, w, h, blocks):
    """blocks: [(text, size, color, bold, space_before)]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for text, size, color, bold, sp in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        if sp:
            p.space_before = Pt(sp)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return tb


def build_page(prs, src_idx, prof, out_text, deck_no, dry=False):
    dst = clone_slide(prs, src_idx)
    new_idx = len(prs.slides) - 1
    # 1) 剥离模板外形状
    keep = TPL[prof]
    for shp in list(dst.shapes):
        key = (round(Emu(shp.left).inches, 2), round(Emu(shp.top).inches, 2))
        if key not in keep:
            shp._element.getparent().remove(shp._element)
    col = C[prof]
    body = clean_lines(out_text)

    if prof == "v1":
        # 页眉标签
        for shp in dst.shapes:
            if shp.has_text_frame and round(Emu(shp.top).inches, 2) == 0.13:
                shp.text_frame.paragraphs[0].runs[0].text = "工单完成效果图"
        add_card(dst, 0.52, 0.30, 1.60, 0.28, col["accent"], col["accent"])
        add_text(dst, 0.52, 0.30, 1.60, 0.28, [("DONE SHOT", 10.5, "#FFFFFF", True, 0)])
        add_text(dst, 0.52, 0.64, 12.29, 0.55,
                 [("工单完成效果图 · 完成后长这样", 25, col["title"], True, 0)])
        add_text(dst, 0.52, 1.16, 12.29, 0.30,
                 [("真实运行 · 参考实现——先看清'交付标准答案'，再动手写自己的", 12, col["muted"], False, 0)])
        add_card(dst, 0.52, 1.50, 12.29, 2.55, col["card"], col["border"])
        add_card(dst, 0.52, 4.15, 12.29, 1.30, col["card2"], col["border"])
        add_card(dst, 0.52, 5.55, 12.29, 1.35, col["card2"], col["border"])
        add_text(dst, 0.82, 1.66, 11.70, 2.25,
                 [("① 参考实现真实运行输出（本工单完成态）", 14, col["head"], True, 0)]
                 + [(l, 10.0, col["body"], False, 1) for l in body])
        add_text(dst, 0.82, 4.30, 11.70, 1.05,
                 [("② 关键自检点（对照工单逐项核对）", 13, col["head"], True, 0),
                  ("□ 代码从头跑到尾不报错（python xx.py 再跑一遍）", 11.5, col["body"], False, 2),
                  ("□ 指标达标：与工单及格线逐项对齐（准确率/F1/耗时）", 11.5, col["body"], False, 1)])
        add_text(dst, 0.82, 5.70, 11.70, 1.10,
                 [("③ 提交与达标（过程性证据）", 13, col["head"], True, 0),
                  ("达标线：D≥60 / C≥70 / B≥80 / A≥90　｜　提交：学号_姓名_工单号 + 运行截图 + 决策记录 + Git≥1次commit", 11, col["body"], False, 2),
                  ("复核：自检 → 互检 → 师查（三层质检）", 11, col["muted"], False, 1)])
    else:
        for shp in dst.shapes:
            if shp.has_text_frame and round(Emu(shp.top).inches, 2) == 0.33:
                shp.text_frame.paragraphs[0].runs[0].text = "工单完成效果图 · 完成后长这样（真实运行 · 参考实现）"
        add_card(dst, 0.83, 1.97, 11.92, 2.50, col["card"], col["border"])
        add_card(dst, 0.83, 4.57, 11.92, 1.15, col["card2"], col["border"])
        add_card(dst, 0.83, 5.82, 11.92, 1.10, col["card2"], col["border"])
        add_text(dst, 1.05, 2.11, 11.45, 2.24,
                 [("① 参考实现真实运行输出（本工单完成态）", 14, col["head"], True, 0)]
                 + [(l, 10.0, col["body"], False, 1) for l in body])
        add_text(dst, 1.05, 4.70, 11.45, 0.95,
                 [("② 关键自检点（对照工单逐项核对）", 13, col["head"], True, 0),
                  ("□ 代码从头跑到尾不报错　□ 指标与工单及格线逐项对齐　□ 随机种子已固定（可复现）", 11.5, col["body"], False, 2)])
        add_text(dst, 1.05, 5.94, 11.45, 0.92,
                 [("③ 提交与达标（过程性证据）", 13, col["accent"], True, 0),
                  ("达标线 D≥60 / C≥70 / B≥80 / A≥90　｜　提交 学号_姓名_工单号 + 运行截图 + 决策记录 + Git≥1次commit", 11, col["body"], False, 2)])
    return dst, new_idx


def do_inject(dry=False, refresh=False):
    data = json.load(open(CACHE, encoding="utf-8"))
    decks = []
    for sub, prof in [("V1授课PPT", "v1"), ("V2授课PPT", "v2")]:
        for p in sorted(glob.glob(os.path.join(PKG, sub, "*.pptx"))):
            if "~$" in p:
                continue
            decks.append((p, prof))
    n_new = n_skip = 0
    for path, prof in decks:
        fn = None
        base = os.path.basename(path)
        for kw, f in KEYMAP:
            if kw in base:
                fn = f
                break
        if fn is None:
            print(f"  [跳过·无匹配] {base}")
            continue
        prs = Presentation(path)
        # 定位全部既有效果图页（refresh 时全部删除，避免重复页残留）
        olds = []
        for i, s in enumerate(prs.slides):
            t = "".join(r.text for sh in s.shapes if sh.has_text_frame
                        for pa in sh.text_frame.paragraphs for r in pa.runs)
            if "工单完成效果图" in t:
                olds.append(i)
        if olds:
            if not refresh:
                n_skip += 1
                continue
            lst = prs.slides._sldIdLst
            for idx in sorted(olds, reverse=True):
                items = list(lst)
                rId = items[idx].get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                prs.part.drop_rel(rId)
                lst.remove(items[idx])
        # 找源模板页 + 插入点
        src_idx = None
        ins_at = None
        for i, s in enumerate(prs.slides):
            t = "".join(r.text for sh in s.shapes if sh.has_text_frame
                        for pa in sh.text_frame.paragraphs for r in pa.runs)
            if prof == "v1" and src_idx is None and t.startswith("工单实战"):
                src_idx, ins_at = i, i + 1
            if prof == "v2" and src_idx is None and t.startswith("实施任务：工单"):
                src_idx, ins_at = i, i + 1
        if src_idx is None:
            print(f"  [跳过·无模板页] {base}")
            continue
        if dry:
            print(f"  [计划] {prof} {base}  源p{src_idx + 1} → 插入p{ins_at + 1}  ({len(prs.slides)}→{len(prs.slides) + 1}页)")
            n_new += 1
            continue
        _, new_idx = build_page(prs, src_idx, prof, data.get(fn, ""), 0)
        move_slide(prs, new_idx, ins_at)
        # 页脚重编
        total = len(prs.slides)
        for i, s in enumerate(prs.slides, 1):
            for sh in s.shapes:
                if not sh.has_text_frame:
                    continue
                for pa in sh.text_frame.paragraphs:
                    for r in pa.runs:
                        if re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", r.text):
                            r.text = f"{i} / {total}"
        prs.save(path)
        n_new += 1
    print(f"\n新建/重建 {n_new} 册 · 幂等跳过 {n_skip} 册")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", action="store_true")
    ap.add_argument("--inject", action="store_true")
    ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()
    if a.run:
        print("=== 运行 05包 12 工单完成态参考实现 ===")
        run_solutions()
    if a.inject or a.dry or a.refresh:
        do_inject(dry=a.dry, refresh=a.refresh)
