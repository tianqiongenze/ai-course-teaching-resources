# -*- coding: utf-8 -*-
"""76号轮·PPT 美工量化审计 v2（校准版）

修正 v1 误报：
  - 栅格判据 → 改为「近失对齐」：同页同带(±0.1in)文本框左边缘差 0.01~0.20in 视为手抖
  - 安全区 → 豁免页眉(T<1.3)与页脚(T>7.0)设计区
  - 重复页 → 豁免「效果图前后双置」配对（标题不同属有意设计）

维度：字体族 / 字号档位 / 品牌色 / 近失对齐 / 越安全区(正文区) / 文本密度 / 真重复页
"""
import glob
import json
import os
import re
from collections import Counter, defaultdict

from pptx import Presentation
from pptx.util import Emu

ROOT = r"D:/Python编程实战/教学实施资料"
BODY_T = (1.3, 7.0)       # 正文区
SAFE_L, SAFE_R = 0.45, 12.95
DENSITY_MAX = 700
NEAR_MISS = (0.01, 0.20)  # 近失对齐区间
DUP_SIM = 0.90
DUP_EXEMPT = ("工单完成效果图", "本工单的工程化项目")


def slide_text(s):
    return "\n".join(sh.text_frame.text for sh in s.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip())


def jac(a, b):
    sa, sb = set(a), set(b)
    return len(sa & sb) / len(sa | sb) if sa and sb else 0.0


def audit_deck(path):
    prs = Presentation(path)
    fonts, sizes, colors = Counter(), Counter(), Counter()
    near_miss, out_safe, dense, dup = [], [], [], []
    texts = []
    for si, s in enumerate(prs.slides, 1):
        t = slide_text(s)
        texts.append(t)
        boxes = []
        for sh in s.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            L, T = Emu(sh.left).inches, Emu(sh.top).inches
            W, H = Emu(sh.width).inches, Emu(sh.height).inches
            for pa in sh.text_frame.paragraphs:
                for r in pa.runs:
                    if r.font.name:
                        fonts[r.font.name] += 1
                    if r.font.size:
                        sizes[round(r.font.size.pt)] += 1
                    try:
                        if r.font.color and r.font.color.type is not None:
                            colors[str(r.font.color.rgb)] += 1
                    except Exception:
                        pass
            if BODY_T[0] <= T <= BODY_T[1]:
                boxes.append((L, T, W, H, sh.text_frame.text[:12]))
                if L < SAFE_L - 0.01 or L + W > SAFE_R + 0.01:
                    out_safe.append((si, round(L, 2), round(L + W, 2), sh.text_frame.text[:14]))
        # 近失对齐：同带内左边缘差 0.01~0.20
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                if abs(boxes[i][1] - boxes[j][1]) <= 0.10:
                    d = abs(boxes[i][0] - boxes[j][0])
                    if NEAR_MISS[0] < d < NEAR_MISS[1]:
                        near_miss.append((si, round(boxes[i][0], 2), round(boxes[j][0], 2),
                                          round(d, 2), boxes[i][4], boxes[j][4]))
        n = len(re.sub(r"\s", "", t))
        if n > DENSITY_MAX:
            dense.append((si, n))
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            if len(texts[i]) > 30 and len(texts[j]) > 30:
                if any(k in texts[i] and k in texts[j] for k in DUP_EXEMPT):
                    continue
                sim = jac(texts[i], texts[j])
                if sim >= DUP_SIM:
                    dup.append((i + 1, j + 1, round(sim, 3)))
    return {
        "pages": len(prs.slides),
        "font_kinds": len(fonts), "fonts": dict(fonts.most_common(4)),
        "size_kinds": len(sizes), "sizes": sorted(sizes),
        "colors": dict(colors.most_common(5)),
        "near_miss_n": len(near_miss), "near_miss": near_miss[:12],
        "out_safe_n": len(out_safe), "out_safe": out_safe[:10],
        "dense": dense, "dup": dup,
    }


def main():
    files = sorted(f for f in glob.glob(os.path.join(ROOT, "**", "*.pptx"), recursive=True)
                   if "~$" not in f and "_docx_work" not in f)
    out, agg = {}, Counter()
    for p in files:
        rel = os.path.relpath(p, ROOT).replace("\\", "/")
        try:
            r = audit_deck(p)
        except Exception as e:
            out[rel] = {"error": str(e)}
            continue
        out[rel] = r
        agg["pages"] += r["pages"]
        if r["font_kinds"] > 2:
            agg["font_mixed"] += 1
        agg["near_miss"] += r["near_miss_n"]
        agg["out_safe"] += r["out_safe_n"]
        agg["dense"] += len(r["dense"])
        agg["dup"] += len(r["dup"])
    json.dump(out, open(os.path.join(ROOT, "_docx_work", "ppt_aesthetic_audit.json"),
                        "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("=" * 78)
    print("PPT 美工审计 v2（校准） · %d 册 / %d 页" % (len(files), agg["pages"]))
    print("=" * 78)
    print("  字体族混用(>2)册 : %d" % agg["font_mixed"])
    print("  近失对齐(0.01~0.2): %d" % agg["near_miss"])
    print("  正文越安全区     : %d" % agg["out_safe"])
    print("  过密页(>700字)   : %d" % agg["dense"])
    print("  真重复页(>0.9)   : %d" % agg["dup"])
    print("-" * 78)
    worst = sorted(((rel, r) for rel, r in out.items() if "error" not in r),
                   key=lambda x: -(x[1]["near_miss_n"] + x[1]["out_safe_n"] * 2))[:5]
    for rel, r in worst:
        print("  %s" % rel[:58])
        print("     近失%d 越界%d 过密%d 重复%d 字号%d档" % (
            r["near_miss_n"], r["out_safe_n"], len(r["dense"]), len(r["dup"]), r["size_kinds"]))
    print("=" * 78)


if __name__ == "__main__":
    main()