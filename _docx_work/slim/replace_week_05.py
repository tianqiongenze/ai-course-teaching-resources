# -*- coding: utf-8 -*-
"""79号轮·05包周次统一：内容替换 v2

映射（教学周 k：k=1→第7周；k≥2→第(k+8)周）
  Wxx（两位，全包统一）:
    W04→W07 W05→W10 W06→W11 W07→W12 W08→W13 W09→W14
    W10→W15 W11→W16 W12→W17 W13→W18 W14→W19 W15→W20
  中文"第N周"——按文件类型用不同旧口径:
    · 工单任务书（旧标注=教学周+4）: 5→7 6→10 7→11 8→12 9→13 10→14 11→15 12→16 13→17 14→18 15→19 16→20
    · 其他（旧学期周=教学周+3）    : 4→7 5→10 6→11 7→12 8→13 9→14 10→15 11→16 12→17 13→18 14→19 15→20
排除 V2授课PPT（已是新口径）。
"""
import glob
import os
import re
import sys

ROOT = r"D:/Python编程实战/教学实施资料/05-AI应用基础48课时版"
WMAP = {"04": "07", "05": "10", "06": "11", "07": "12", "08": "13", "09": "14",
        "10": "15", "11": "16", "12": "17", "13": "18", "14": "19", "15": "20"}
GONGDAN = {"5": "7", "6": "10", "7": "11", "8": "12", "9": "13", "10": "14",
           "11": "15", "12": "16", "13": "17", "14": "18", "15": "19", "16": "20"}
SEM = {"4": "7", "5": "10", "6": "11", "7": "12", "8": "13", "9": "14",
       "10": "15", "11": "16", "12": "17", "13": "18", "14": "19", "15": "20"}
TAG = "\u0001"


RANGE_MAP = [
    ("第4-15周", "第7周、第10-20周"),
    ("第4-15 周", "第7周、第10-20周"),
    ("第4-5周", "第7周、第10周"),
    ("第4-5 周", "第7周、第10周"),
    ("第6-9周", "第11-14周"),
    ("第6-9 周", "第11-14周"),
    ("第10-11周", "第15-16周"),
    ("第10-11 周", "第15-16周"),
    ("第12-13周", "第17-18周"),
    ("第12-13 周", "第17-18周"),
    ("W4-11", "第7周、第10-16周"),
    ("W4-15", "第7周、第10-20周"),
]


def repl(t, cmap):
    n = 0
    for a, b in RANGE_MAP:
        if a in t:
            n += t.count(a)
            t = t.replace(a, b)

    def wf(m):
        nonlocal n
        k = m.group(1)
        if k in WMAP:
            n += 1
            return TAG + "W" + WMAP[k]
        return m.group(0)
    t = re.sub(r"W(\d{2})", wf, t)

    def cf(m):
        nonlocal n
        k = m.group(1)
        if k in cmap:
            n += 1
            return "第" + TAG + cmap[k] + "周"
        return m.group(0)
    t = re.sub(r"第(\d{1,2})周", cf, t)
    return t.replace(TAG, ""), n


def main():
    dry = "--apply" not in sys.argv
    os.chdir(ROOT)
    jobs = []
    # V1 授课 PPT（排除 V2）
    for p in glob.glob("V1授课PPT/*.pptx"):
        jobs.append((p, "pptx", SEM))
    # README 与其他根级 md（学期周口径）
    for p in glob.glob("*.md"):
        jobs.append((p, "md", SEM))
    # 工单任务书（工单口径）
    for p in glob.glob("工单任务书/*.md"):
        jobs.append((p, "md", GONGDAN))
    # 分层工单 docx（工单口径）
    for p in glob.glob("分层工单/*.docx"):
        jobs.append((p, "docx", GONGDAN))
    # 根级 docx（期末项目册等，学期周口径）
    for p in glob.glob("*.docx"):
        jobs.append((p, "docx", SEM))

    tot = 0
    for path, kind, cmap in jobs:
        if os.path.basename(path).startswith("~$"):
            continue
        try:
            if kind == "md":
                s = open(path, encoding="utf-8").read()
                nt, c = repl(s, cmap)
                if c and not dry:
                    open(path, "w", encoding="utf-8", newline="\n").write(nt)
            elif kind == "pptx":
                c = do_pptx(path, cmap, dry)
            else:
                c = do_docx(path, cmap, dry)
        except Exception as e:
            print("  ERR %s: %s" % (path[:50], e))
            continue
        if c:
            print("  %-56s %d 处" % (path[:56], c))
            tot += c
    print("%s：共 %d 处" % ("空跑" if dry else "已替换", tot))


def do_pptx(path, cmap, dry):
    from pptx import Presentation
    prs = Presentation(path)
    total = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for pa in sh.text_frame.paragraphs:
                    for r in pa.runs:
                        nt, c = repl(r.text, cmap)
                        if c:
                            r.text = nt
                            total += c
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for pa in cell.text_frame.paragraphs:
                            for r in pa.runs:
                                nt, c = repl(r.text, cmap)
                                if c:
                                    r.text = nt
                                    total += c
    if total and not dry:
        prs.save(path)
    return total


def do_docx(path, cmap, dry):
    from docx import Document
    doc = Document(path)
    total = 0

    def walk(paras):
        nonlocal total
        for p in paras:
            for r in p.runs:
                nt, c = repl(r.text, cmap)
                if c:
                    r.text = nt
                    total += c
    walk(doc.paragraphs)
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                walk(cell.paragraphs)
    if total and not dry:
        doc.save(path)
    return total


if __name__ == "__main__":
    main()
