# -*- coding: utf-8 -*-
"""78号轮·《AI应用基础（48课时版）》课程整体说课 PPT 生成器（30 页）

依据：课程标准（B46010304X · 48学时 · 第4-15周）+ 授课计划 + 05包全部资料
考核方案（本轮口径）：平时40（工单10次）+ 课堂表现10 + 小组项目30 + 期末考试20 = 100
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, "说课PPT", "AI应用基础（48课时版）-课程整体说课（30页·2026版）.pptx")

# ── 品牌色 ──
NAVY = RGBColor(0x01, 0x03, 0x28)      # 主色·深蓝
TEAL = RGBColor(0x00, 0xA6, 0xA6)      # 强调·蓝绿
ORANGE = RGBColor(0xF5, 0x67, 0x0F)    # 点缀·智造橙
PURPLE = RGBColor(0x23, 0x1D, 0x3D)
GRAY = RGBColor(0x44, 0x54, 0x6A)
LGRAY = RGBColor(0x8A, 0x97, 0xB5)
BG = RGBColor(0xF7, 0xF8, 0xFA)
CARD = RGBColor(0xEE, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"
FOOT = "重庆两江智造集团 · AI应用基础（48课时版）"

SECTIONS = ["课程定位", "学情分析", "课程目标", "内容与结构",
            "教学设计", "教学资源", "考核评价", "特色与成效"]


def blank(prs, bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def txt(s, l, t, w, h, text, size=14, bold=False, color=NAVY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.15):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def box(s, l, t, w, h, fill=CARD, line=None, radius=True, lw=1.0):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def footer(s, page, total=30):
    box(s, 0.55, 7.02, 12.23, 0.02, fill=RGBColor(0xD8, 0xDE, 0xE8), radius=False)
    txt(s, 0.6, 7.08, 8.0, 0.3, FOOT, 10, False, LGRAY)
    txt(s, 11.0, 7.08, 1.8, 0.3, "%d / %d" % (page, total), 10, False, LGRAY, PP_ALIGN.RIGHT)


def head(s, title, sub="", page=0):
    box(s, 0.55, 0.42, 0.13, 0.62, fill=TEAL, radius=False)
    txt(s, 0.85, 0.42, 10.5, 0.45, title, 22, True, NAVY)
    if sub:
        txt(s, 0.85, 0.88, 11.5, 0.34, sub, 12, False, GRAY)
    if page:
        footer(s, page)


def divider(prs, no, name, sub, page):
    s = blank(prs, NAVY)
    txt(s, 0.9, 2.15, 3.0, 1.6, no, 96, True, RGBColor(0x1F, 0x3A, 0x5C))
    box(s, 3.5, 2.62, 0.06, 1.0, fill=ORANGE, radius=False)
    txt(s, 3.85, 2.62, 8.5, 0.9, name, 40, True, WHITE)
    txt(s, 3.9, 3.62, 8.5, 0.5, sub, 15, False, RGBColor(0x8A, 0x97, 0xB5))
    # 右侧环节导航
    txt(s, 10.4, 0.6, 2.4, 6.3, "\n".join(
        ("▶ " if x == name else "   ") + x for x in SECTIONS), 12, False,
        RGBColor(0x6C, 0x7C, 0x9C), PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, 1.9)
    txt(s, 0.9, 7.08, 8.0, 0.3, FOOT, 10, False, RGBColor(0x4A, 0x58, 0x78))
    txt(s, 11.0, 7.08, 1.8, 0.3, "%d / 30" % page, 10, False, RGBColor(0x4A, 0x58, 0x78), PP_ALIGN.RIGHT)


def cards(s, items, top=1.5, height=4.6, gap=0.28, left=0.6, right=12.75, fs=12):
    """items: [(标题, 正文)]  横向等分卡片"""
    n = len(items)
    total = right - left
    w = (total - gap * (n - 1)) / n
    for i, (t, b) in enumerate(items):
        l = left + i * (w + gap)
        box(s, l, top, w, height, fill=WHITE, line=RGBColor(0xD8, 0xDE, 0xE8))
        box(s, l, top, w, 0.52, fill=TEAL, radius=True)
        txt(s, l + 0.1, top + 0.04, w - 0.2, 0.44, t, 14, True, WHITE,
            PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(s, l + 0.2, top + 0.72, w - 0.4, height - 0.9, b, fs, False, GRAY,
            PP_ALIGN.LEFT, MSO_ANCHOR.TOP, 1.3)


def table(s, data, top=1.5, left=0.6, width=12.15, height=4.6,
          col_w=None, fs=11, head_fs=12, row_h=None):
    rows, cols = len(data), len(data[0])
    shp = s.shapes.add_table(rows, cols, Inches(left), Inches(top),
                             Inches(width), Inches(height))
    tb = shp.table
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            tb.columns[j].width = Emu(int(Inches(width) * cw / tot))
    if row_h:
        for i, rh in enumerate(row_h):
            tb.rows[i].height = Inches(rh)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.text = ""
            c.margin_left = Inches(0.08)
            c.margin_right = Inches(0.06)
            c.margin_top = Inches(0.03)
            c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else CARD)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or cols <= 3 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(head_fs if i == 0 else fs)
            r.font.bold = (i == 0)
            r.font.color.rgb = WHITE if i == 0 else NAVY
            r.font.name = FONT
    return tb


def steps(s, items, top=2.4, left=0.75, right=12.6, h=1.5):
    n = len(items)
    gap = 0.42
    w = (right - left - gap * (n - 1)) / n
    for i, (t, b) in enumerate(items):
        l = left + i * (w + gap)
        box(s, l, top, w, h, fill=WHITE, line=TEAL)
        txt(s, l + 0.06, top + 0.12, w - 0.12, 0.4, t, 14, True, TEAL,
            PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(s, l + 0.08, top + 0.6, w - 0.16, h - 0.7, b, 11, False, GRAY,
            PP_ALIGN.CENTER, MSO_ANCHOR.TOP, 1.25)
        if i < n - 1:
            txt(s, l + w + 0.02, top + 0.5, gap - 0.04, 0.5, "▶", 14, True, ORANGE,
                PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def stat(s, l, t, w, h, num, label, color=TEAL):
    box(s, l, t, w, h, fill=WHITE, line=RGBColor(0xD8, 0xDE, 0xE8))
    txt(s, l, t + 0.16, w, 0.9, num, 40, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, l, t + 1.05, w, 0.5, label, 12, False, GRAY, PP_ALIGN.CENTER, MSO_ANCHOR.TOP, 1.2)
