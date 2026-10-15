# -*- coding: utf-8 -*-
"""效果图注入（72号轮）——用真实参考实现运行输出替换"（教师课后替换）"空白占位页
用法: python inject_shots.py --run    (先跑解，缓存)   |  python inject_shots.py --inject
"""
import os, sys, re, glob, io, json, contextlib, copy, argparse
from pptx import Presentation
from pptx.util import Emu, Pt, Inches

ROOT = r"D:/Python编程实战/教学实施资料"
CACHE = os.path.join(ROOT, "_docx_work", "slim", "_shots.json")

# 02包: 周 -> solutions 函数
PY_MAP = {"W01":"w01","W02":"w02","W03":"w03","W04":"w04","W05":"w05","W06":"w06",
          "W07":"w07","W08":"w08","W09":"w09","W10":"w10","W11":"w11","W12":"w12"}
# 01包: 次课序号 -> solutions 函数
AI_MAP = {"W04":"m11","W05":"m12","W06":"m21","W07":"m22","W08":"m31","W09":"m32","W10":"m41","W11":"m42"}

def run_solutions(pkg_dir, ref_sub, funcs):
    """在 pkg_dir 下运行 ref_sub/solutions.py 的 funcs，返回 {func: stdout}"""
    sys.path.insert(0, os.path.join(pkg_dir, ref_sub))
    for m in list(sys.modules):
        if m == "solutions": del sys.modules[m]
    cwd = os.getcwd(); os.chdir(pkg_dir)
    out = {}
    try:
        import solutions
        for f in funcs:
            buf = io.StringIO()
            try:
                with contextlib.redirect_stdout(buf):
                    getattr(solutions, f)()
                out[f] = buf.getvalue()
            except Exception as e:
                out[f] = f"__ERR__ {type(e).__name__}: {str(e)[:80]}"
    finally:
        os.chdir(cwd); sys.path.pop(0)
    return out

def do_run():
    os.environ["NO_PROXY"] = "127.0.0.1,localhost"
    os.environ["no_proxy"] = "127.0.0.1,localhost"
    # 进程内启动 mock LLM 服务（供 01包 m42 使用），避免依赖外部后台进程
    try:
        import threading, time as _t, uvicorn
        sys.path.insert(0, os.path.join(ROOT, "01-AI应用基础", "数据与工具"))
        cwd = os.getcwd(); os.chdir(os.path.join(ROOT, "01-AI应用基础", "数据与工具"))
        from mock_llm_server import app as _llm_app
        os.chdir(cwd)
        cfg = uvicorn.Config(_llm_app, host="127.0.0.1", port=8000, log_level="error")
        srv = uvicorn.Server(cfg)
        threading.Thread(target=srv.run, daemon=True).start()
        _t.sleep(4)
        print("  [llm] mock LLM 服务已就绪 127.0.0.1:8000")
    except Exception as e:
        print("  [llm] 启动失败:", e)
    res = {}
    py_dir = os.path.join(ROOT, "02-程序设计基础")
    res["py"] = run_solutions(py_dir, "教师参考", list(PY_MAP.values()))
    ai_dir = os.path.join(ROOT, "01-AI应用基础")
    res["ai"] = run_solutions(ai_dir, "教师参考", list(AI_MAP.values()))
    json.dump(res, open(CACHE,"w",encoding="utf-8"), ensure_ascii=False, indent=1)
    for k,v in res.items():
        for f,t in v.items():
            print(f"  [{k}] {f}: {'ERR' if t.startswith('__ERR__') else str(len(t.splitlines()))+'行'} {t.splitlines()[0][:60] if t.strip() else ''}")

KEYMARK = ("结论", "要点", "发现", "路径", "合计", "召回", "达标", "目标", "说明", "关键", "分类", "指标")


def clean_lines(text, maxlines=8, maxw=40):
    ls = [l.rstrip() for l in text.splitlines() if l.strip()]
    ls = [l if len(l) <= maxw else l[:maxw - 1] + "…" for l in ls]
    # 折叠连续同类行（数字归一后相同）
    merged = []
    for l in ls:
        key = re.sub(r"\d+", "#", l)
        if merged and re.sub(r"\d+", "#", merged[-1][0]) == key:
            merged[-1][1] += 1
            continue
        merged.append([l, 1])
    res = [l + (f"  ×{c}" if c > 1 else "") for l, c in merged]
    # 超长时按"关键行优先、其余按原序补齐"裁剪（保持原文顺序，不重排）
    if len(res) > maxlines:
        first, body = res[0], res[1:]
        keys_idx = [i for i, l in enumerate(body) if any(k in l for k in KEYMARK)]
        others_idx = [i for i in range(len(body)) if i not in keys_idx]
        keep = set(keys_idx[:maxlines - 1])
        for i in others_idx:
            if len(keep) >= maxlines - 1:
                break
            keep.add(i)
        res = [first] + [body[i] for i in sorted(keep)][:maxlines - 1]
    return res

def set_frame(tb, lines, size, card=None, color=None, bold_first=False, left_align=True):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    if card is not None:
        tb.left = Emu(int(card.left) + Emu(int(0.18*914400)))
        tb.top = Emu(int(card.top) + Emu(int(0.14*914400)))
        tb.width = Emu(int(card.width) - Emu(int(0.36*914400)))
        tb.height = Emu(int(card.height) - Emu(int(0.28*914400)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    base_p = tf.paragraphs[0]
    if not base_p.runs:
        base_p.add_run()
    proto = copy.deepcopy(base_p._p)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    r0 = tf.paragraphs[0].runs[0]
    r0.text = lines[0] if lines else ""
    for r in list(tf.paragraphs[0].runs)[1:]:
        r._r.getparent().remove(r._r)
    for r in tf.paragraphs[0].runs:
        r.font.size = Pt(size); r.font.bold = bold_first
        if color is not None: r.font.color.rgb = color
    for ln in lines[1:]:
        newp = copy.deepcopy(proto)
        tf._txBody.append(newp)
    for i, p in enumerate(list(tf.paragraphs)[1:], 1):
        for r in list(p.runs)[1:]:
            r._r.getparent().remove(r._r)
        if not p.runs: p.add_run()
        p.runs[0].text = lines[i] if i < len(lines) else ""
        p.runs[0].font.size = Pt(size); p.runs[0].font.bold = False
        if color is not None: p.runs[0].font.color.rgb = color
    for p in tf.paragraphs:
        p.space_before = Pt(0); p.space_after = Pt(0)
        if left_align:
            p.alignment = PP_ALIGN.LEFT


def relayout_4to3(cards, caps):
    """把 2×2 四卡重排为『宽输出卡 + 两小卡 + 宽提交卡』（利用下方闲置空间）"""
    if len(cards) != 4 or len(caps) != 4:
        return
    pos = [(0.77, 2.02, 11.79, 2.62), (0.77, 4.72, 5.03, 1.10),
           (7.53, 4.72, 5.03, 1.10), (0.77, 5.92, 11.79, 1.10)]
    for shp, (x, y, w, h) in zip(cards, pos):
        shp.left, shp.top, shp.width, shp.height = (
            Inches(x), Inches(y), Inches(w), Inches(h))
    for cap, (x, y, w, h) in zip(caps, pos):
        cap.left, cap.top = Inches(x + 0.18), Inches(y + 0.14)
        cap.width, cap.height = Inches(w - 0.36), Inches(h - 0.28)

def do_inject(refresh=False):
    data = json.load(open(CACHE, encoding="utf-8"))
    jobs = []
    for p in sorted(glob.glob(os.path.join(ROOT,"02-程序设计基础","**","*.pptx"), recursive=True)):
        if "~$" in p: continue
        m = re.search(r"(W\d{2})-", os.path.basename(p))
        if m: jobs.append((p, "py", PY_MAP[m.group(1)]))
    for p in sorted(glob.glob(os.path.join(ROOT,"01-AI应用基础","**","*.pptx"), recursive=True)):
        if "~$" in p: continue
        m = re.search(r"(W\d{2})-", os.path.basename(p))
        if m and m.group(1) in AI_MAP: jobs.append((p, "ai", AI_MAP[m.group(1)]))
    n=0; nref=0
    for path, grp, fn in jobs:
        out = data[grp].get(fn,"")
        prs = Presentation(path)
        target=None
        for s in prs.slides:
            t="".join(r.text for sh in s.shapes if sh.has_text_frame for pa in sh.text_frame.paragraphs for r in pa.runs)
            if ("课堂实况与运行截图" in t) or ("工单完成效果图" in t): target=s; break
        if target is None: continue
        has_new = any(("工单完成效果图" in r.text) for sh in target.shapes if sh.has_text_frame
                      for pa in sh.text_frame.paragraphs for r in pa.runs)
        if has_new and not refresh:
            continue
        cards=[]; caps=[]; title=None
        for sh in target.shapes:
            if not sh.has_text_frame: continue
            tx=sh.text_frame.text.strip()
            if tx.startswith("课堂实况与运行截图") or tx.startswith("工单完成效果图"): title=sh
            elif tx.startswith("［插入运行截图") or tx.startswith(("①","②","③","④")): caps.append(sh)
        for sh in target.shapes:
            if sh.shape_type == 1 and 1.9 <= Emu(sh.top).inches <= 6.5 \
                    and Emu(sh.width).inches >= 4.5 and 0.5 <= Emu(sh.height).inches <= 3.0:
                cards.append(sh)
        cards.sort(key=lambda sh:(round(Emu(sh.top).inches,1), round(Emu(sh.left).inches,1)))
        if title:
            for pa in title.text_frame.paragraphs:
                for r in pa.runs: r.text = ""
            title.text_frame.paragraphs[0].runs[0].text = "工单完成效果图 · 完成后长这样（真实运行 · 参考实现）"
        body = clean_lines(out, maxlines=12, maxw=64)
        f2 = ["② 关键自检点（对照工单逐项核对）","□ 代码从头跑到尾不报错","□ 输出格式符合工单样例（对齐/单位/小数位）","□ 文件名 学号_姓名_工单号.py"]
        f3 = ["③ 四级达标线（当堂评分）","D·学徒 ≥60 跟做完成核心步骤","C·操作员 ≥70 基础功能完成","B·组长 ≥80 全功能+规范+帮带"]
        f4 = ["④ 提交规范（过程性证据）","□ 运行截图贴工单过程记录栏","□ Git 本次课≥1条 commit","□ 互查签字后提交"]
        frames=[["① 参考实现真实运行输出（本课工单·完成态）"]+body, f2, f3, f4]
        caps.sort(key=lambda sh:(round(Emu(sh.top).inches,1), round(Emu(sh.left).inches,1)))
        if len(cards) == 4:
            relayout_4to3(cards, caps)
        for i,(cap, lines) in enumerate(zip(caps, frames)):
            set_frame(cap, lines, 9.5 if i==0 else 10.0, card=cards[i] if i < len(cards) else None)
        prs.save(path); n+=1
        if has_new: nref+=1
    print(f"已注入/刷新 {n} 册（其中刷新 {nref} 册）")

if __name__=="__main__":
    ap=argparse.ArgumentParser(); ap.add_argument("--run",action="store_true")
    ap.add_argument("--inject",action="store_true"); ap.add_argument("--refresh",action="store_true")
    a=ap.parse_args()
    if a.run: do_run()
    if a.inject or a.refresh: do_inject(refresh=a.refresh)
